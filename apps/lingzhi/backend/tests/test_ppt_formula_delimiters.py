"""Source math delimiters must not invalidate an otherwise renderable page."""
from copy import deepcopy

import pytest
from pptx import Presentation
from pptx.util import Pt

from ppt_formula_projection import project_formula
from ppt_native_scene import audit_scene, render_scene
from teacher_script_ppt import validate_block_pages
from backend.tests.test_teacher_script_ppt import sample


@pytest.mark.parametrize("source", [r"\(p = 1/2\)", r"\[p = 1/2\]", "$p = 1/2$", "$$p = 1/2$$"])
def test_standalone_formula_accepts_matched_math_delimiters(source):
    assert project_formula(source) == "p = 1/2"


@pytest.mark.parametrize("source", [r"\(p = 1/2\]", r"\(p = 1/2", r"\(\unknown{x}\)"])
def test_formula_delimiters_do_not_hide_invalid_notation(source):
    with pytest.raises(ValueError, match="teaching_formula_not_supported"):
        project_formula(source)


@pytest.mark.parametrize("wrapper", [(r"\(", r"\)"), (r"\[", r"\]"), ("$$", "$$")])
def test_wrapped_matrix_retains_measured_rows(wrapper):
    matrix = r"\begin{bmatrix}1&2\\3&4\end{bmatrix}"
    assert project_formula(wrapper[0] + matrix + wrapper[1]) == project_formula(matrix)


def test_saved_formula_page_validates_and_exports_without_rewriting_source(tmp_path):
    from course_document import CourseBlock, CourseDocument, CourseSection, stable_hash
    from ppt_teaching_manuscript import compile_teaching_manuscript
    from teacher_script_ppt import lower_bundle_pages

    template, contract, _ = sample()
    source = r"每个节点以概率 \(p = 1/2\) 拥有更高层级的指针。"
    formula = r"\(p = 1/2\)"
    page = {"layout_id": template.layout_id("formula"), "page_goal": "说明随机层级的概率", "fields": {
        "title": "跳表的随机参数", "notes": "根据来源解释层级概率。",
        "formula": {"sources": [{"block_id": "b", "quote": formula}]},
        "explanation": {"text": "节点以概率 p = 1/2 拥有更高层级指针", "sources": [{"block_id": "b", "quote": source}]},
    }}
    block = {**contract["modules"][0], "content": source, "ppt_pages": [page]}
    before = deepcopy(block)
    validate_block_pages(block, template)
    document = CourseDocument(course_id="formula-test", title="跳表", document_revision="doc",
        sections=[CourseSection(section_id="s", title="跳表", position=0)],
        blocks=[CourseBlock(block_id="b", section_id="s", position=0, payload={"markdown": source},
                            internal_revision=stable_hash(source, prefix="block_"))])
    graph, planned = lower_bundle_pages(document, template, [block])
    manuscript = compile_teaching_manuscript(document, graph, template, {}, planned)
    scene = manuscript.pages[0].resolved_scenes[0]
    obj = next(obj for obj in scene.objects if obj.element_id == "formula")
    assert obj.text == "p = 1/2"
    deck = Presentation()
    deck.slide_width, deck.slide_height = Pt(960), Pt(540)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    render_scene(slide, scene)
    path = tmp_path / "formula.pptx"
    deck.save(path)
    assert audit_scene(Presentation(path).slides[0], scene)["passed"]
    assert block == before
