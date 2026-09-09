"""Regression cases for the handout -> retry -> saved manuscript boundary."""
import asyncio
import json
import time
from copy import deepcopy

import pytest

from backend.tests.test_teacher_script_ppt import sample, TEXT
from backend.tests.test_teacher_ppt_joint_workflow import workflow, generate
from routers import teacher_lesson_authoring as routes
from teacher_script_ppt import CONTRACT, generate_bundle, validate_block_pages


def unity_page(template, source):
    return {"layout_id": template.layout_id("bullets"), "page_goal": "查看运行结果",
            "fields": {"title": "运行结果", "notes": "查看输出", "points": [
                {"text": "在 Console 查看 Start 和 Update", "sources": source}]}}


@pytest.mark.parametrize("source", [
    [{"block_id": "b", "quote": "创建 1 个 Cube。"}],
    [{"quote_id": "q_unknown"}],
    ["q_unknown"],
])
def test_rebinds_multiple_source_claims_before_spending_repair_budget(source):
    template, contract, _ = sample()
    text = "创建 1 个 Cube。Start 在启动时执行。Update 逐帧执行。Console 显示输出。"
    page = unity_page(template, source)
    seed = {**contract["modules"][0], "content": text, "ppt_pages": [page],
            "ppt_repair_attempts": [2], "generation_contract_version": CONTRACT}

    async def unexpected(*args, **kwargs):
        pytest.fail("source binding should use the frozen handout")

    result = asyncio.run(generate_bundle(invoke=unexpected, contract=contract, instructions="",
        template=template, seed_blocks={"b": seed}, immutable_handout=True))
    block = result["blocks"][0]
    assert not block["ppt_errors"]
    assert block["ppt_pages"][0]["fields"]["points"][0]["text"] == page["fields"]["points"][0]["text"]
    validate_block_pages(block, template)


def test_unknown_layout_enters_bounded_page_repair():
    template, contract, page = sample()
    bad = {**deepcopy(page), "layout_id": "nonexistent-layout"}
    calls = []

    async def invoke(*args, **kwargs):
        calls.append(args)
        return json.dumps({"pages": [page]})

    seed = {**contract["modules"][0], "content": TEXT, "ppt_pages": [bad], "generation_contract_version": CONTRACT}
    result = asyncio.run(generate_bundle(invoke=invoke, contract=contract, instructions="",
        template=template, seed_blocks={"b": seed}, immutable_handout=True))
    assert len(calls) == 1
    validate_block_pages(result["blocks"][0], template)


def _wait(client, job_id):
    for _ in range(500):
        job = client.get(f"/api/teacher/courses/course-1/lesson-jobs/{job_id}").json()["job"]
        if job["status"] not in {"pending", "running"}:
            return job
        time.sleep(.01)
    pytest.fail("PPT completion did not finish")


def test_http_retry_uses_latest_pages_and_fresh_budget_from_checkpoint(workflow, monkeypatch):
    client, repository, _ = workflow
    generate(client, complete_ppt=False)
    lesson = repository.lesson("course-1", "L1-1")
    sections = deepcopy(lesson["script_revisions"][-1]["sections"])
    old = sections[0]["blocks"][0]
    old.update(ppt_pages=[{"old": True}], ppt_page_groups=[[{"old": True}]],
               ppt_repair_attempts=[2], generation_contract_version=CONTRACT)
    lesson = repository.save_script_revision("course-1", "L1-1", sections,
        source_lesson_plan_revision_id=lesson["working_revision_id"], generation_source="teacher_edit")
    template = sample()[0]
    job = repository.create_job("course-1", "L1-1", job_type="teacher_lesson_ppt_manuscript_generation", request_id="retry")
    repository.bind_ppt_completion("course-1", "L1-1", job["id"], lesson["working_script_revision_id"],
        lesson["working_revision_id"], template.model_dump(mode="json"))
    checkpoint = {**deepcopy(old), "ppt_pages": [{"new": True}], "ppt_page_groups": [[{"new": True}]]}
    repository.save_script_bundle_checkpoint("course-1", job["id"], checkpoint)
    repository.update_job("course-1", job["id"], status="failed")
    tm = client.app.dependency_overrides[routes.require_task_manager]()
    original = tm.course_service.generate_teacher_script_section
    captured = []

    async def capture(**kwargs):
        captured.append(deepcopy(kwargs["bundle_seed_blocks"]))
        # Run the real generation service on its normal empty initial input.
        kwargs["bundle_seed_blocks"] = {key: {k: v for k, v in b.items() if not k.startswith("ppt_")}
                                        for key, b in kwargs["bundle_seed_blocks"].items()}
        return await original(**kwargs)

    tm.course_service.generate_teacher_script_section = capture
    source = routes._teacher_v6_source

    def legacy_source(*args):
        values = list(source(*args))
        # Historical joint-generation revisions retained page repair fields.
        revision = next(r for r in values[3]["script_revisions"]
                        if r["revision_id"] == lesson["working_script_revision_id"])
        revision["sections"][0]["blocks"][0].update({k: deepcopy(v) for k, v in old.items() if k.startswith("ppt_")})
        return tuple(values)

    monkeypatch.setattr(routes, "_teacher_v6_source", legacy_source)
    response = client.post("/api/teacher/courses/course-1/lessons/L1-1/ppt-v6/manuscript/complete",
        json={"source_script_revision_id": lesson["working_script_revision_id"], "task_id": job["id"]})
    assert response.status_code == 202, response.text
    assert _wait(client, job["id"])["status"] == "completed"
    seed = captured[0][old["block_id"]]
    assert seed["ppt_pages"] == checkpoint["ppt_pages"]
    assert seed["ppt_page_groups"] == checkpoint["ppt_page_groups"]
    assert "ppt_repair_attempts" not in seed


def test_page_errors_fail_before_final_compilation(workflow):
    client, repository, _ = workflow
    generate(client, complete_ppt=False)
    tm = client.app.dependency_overrides[routes.require_task_manager]()
    original = tm.course_service.generate_teacher_script_section
    phases = []
    update = repository.update_job

    def record(*args, **kwargs):
        phases.append(kwargs.get("phase"))
        return update(*args, **kwargs)

    repository.update_job = record

    async def invalid(**kwargs):
        result = await original(**kwargs)
        block = result["blocks"][0]
        block["ppt_errors"] = [{"block_id": block["block_id"], "page_index": 0,
                                "message": "teaching_fact_token_unsupported:item-1"}]
        return result

    tm.course_service.generate_teacher_script_section = invalid
    lesson = repository.lesson("course-1", "L1-1")
    response = client.post("/api/teacher/courses/course-1/lessons/L1-1/ppt-v6/manuscript/complete",
        json={"source_script_revision_id": lesson["working_script_revision_id"]})
    assert response.status_code == 202
    job = _wait(client, response.json()["job"]["id"])
    assert job["status"] == "failed"
    assert "ppt_manuscript_compiling" not in phases
    assert "ppt_manuscript_saving" not in phases
    assert job["error"]["failed_step"] == "sources"


def test_checkpoint_checks_handout_revision_in_every_progress_phase(workflow, monkeypatch):
    client, repository, _ = workflow
    generate(client, complete_ppt=False)
    lesson = repository.lesson("course-1", "L1-1")
    job = repository.create_job("course-1", "L1-1", job_type="teacher_lesson_ppt_manuscript_generation", request_id="freshness")
    repository.bind_ppt_completion("course-1", "L1-1", job["id"], lesson["working_script_revision_id"],
        lesson["working_revision_id"], sample()[0].model_dump(mode="json"))
    repository.update_job("course-1", job["id"], phase="ppt_page_validation")
    value = repository.load("course-1")
    value["lessons"]["L1-1"]["working_script_revision_id"] = "changed-in-another-window"
    monkeypatch.setattr(repository, "load", lambda *_: deepcopy(value))
    from teacher_lesson_authoring import TeacherLessonAuthoringError
    with pytest.raises(TeacherLessonAuthoringError, match="讲义已变化"):
        repository.save_script_bundle_checkpoint("course-1", job["id"], lesson["script_revisions"][-1]["sections"][0]["blocks"][0])


@pytest.mark.parametrize("multiline_scene", [False, True])
def test_saved_manuscript_build_stream_and_download_keep_all_pages_and_notes(workflow, monkeypatch, tmp_path, multiline_scene):
    from io import BytesIO
    from pptx import Presentation
    from teaching_representations import TeachingRepresentationRepository

    client, repository, calls = workflow
    monkeypatch.setattr(routes, "teaching_representation_repository", TeachingRepresentationRepository(tmp_path / "representations"))
    monkeypatch.setattr(routes, "_capture_generation_source_snapshot", lambda **_: None)

    async def forbidden(*args, **kwargs):
        pytest.fail("Export must not rewrite the saved manuscript through a model")

    monkeypatch.setattr(routes, "build_ai_base_story_planner_v6", lambda: forbidden)
    monkeypatch.setattr(routes, "build_ai_base_visual_planner_v2", lambda: forbidden)
    if multiline_scene:
        tm = client.app.dependency_overrides[routes.require_task_manager]()
        original = tm.course_service.generate_teacher_script_section

        async def with_multiline_bullets(**kwargs):
            result = await original(**kwargs)
            if kwargs.get("immutable_handout"):
                block = result["blocks"][0]
                block["ppt_pages"][0] = {
                    "layout_id": sample()[0].layout_id("bullets"), "page_goal": "比较执行方式",
                    "fields": {"title": "执行方式和任务依赖关系的比较" * 2, "notes": "检查任务依赖",
                        "points": [{"text": text, "sources": [{"block_id": block["block_id"], "quote": TEXT}]}
                                   for text in (TEXT[:48], "独立任务可以并行", "存在依赖的任务需按顺序执行")]},
                }
            return result

        tm.course_service.generate_teacher_script_section = with_multiline_bullets
    generate(client)
    state = repository.current_v6_ppt_manuscript("course-1", "L1-1")
    base = "/api/teacher/courses/course-1/lessons/L1-1/ppt-v6"
    response = client.post(base + "/build/stream", json={"expected_manuscript_revision": state["revision"]})
    assert response.status_code == 200, response.text
    events = [json.loads(line[6:]) for line in response.text.splitlines() if line.startswith("data: ")]
    finished = next((e for e in events if e.get("event") == "build_complete"), None)
    assert finished is not None, events
    assert finished["job"]["status"] == "completed"
    representation = finished["build"]["representation_id"]
    output = client.get(base + f"/{representation}/export.pptx")
    assert output.status_code == 200, output.text[:300] if output.status_code != 200 else ""
    assert output.content.startswith(b"PK")
    slides = Presentation(BytesIO(output.content)).slides
    assert len(slides) == state["manuscript"]["page_count"]
    for slide in slides:
        assert TEXT in slide.notes_slide.notes_text_frame.text
    assert len(calls) == 2
    assert repository.current_v6_ppt_manuscript("course-1", "L1-1")["revision"] == state["revision"]


def test_confirmed_export_stops_after_three_failed_attempts(monkeypatch):
    from types import SimpleNamespace
    from slide_deck_v6_models import V6BuildError
    job = {"id": "export", "type": "teacher_lesson_ppt_generation", "status": "running", "attempt_number": 0}
    calls = []

    def update(*args, **changes):
        job.update(changes)
        return deepcopy(job)

    async def build(**kwargs):
        calls.append(kwargs)
        if len(calls) > 3:
            pytest.fail("confirmed export kept retrying after three failures")
        raise V6BuildError(stage="export", code="export_runtime_failed", message="renderer unavailable", retryable=True)

    async def no_wait(*args):
        job["next_retry_at"] = None

    monkeypatch.setattr(routes.asyncio, "sleep", no_wait)
    repo = SimpleNamespace(get_job=lambda *_: deepcopy(job), update_job=update)
    with pytest.raises(routes._TeacherPptV6JobStopped):
        asyncio.run(routes._build_teacher_ppt_attempt(repository=repo, course_id="course", task_id="export",
            orchestrator=SimpleNamespace(build=build), confirmed_manuscript=object(), source_revision_provider=lambda: "doc"))
    assert len(calls) == 3
    assert job["status"] == "paused"
    assert job["error"]["code"] == "export_runtime_failed"


def test_stopped_export_event_updates_client_job_status():
    job = {"id": "export", "status": "paused", "phase": "paused", "message": "renderer unavailable"}
    event = routes._teacher_ppt_stopped_event(job)
    assert event["job"]["status"] == "paused"
