from copy import deepcopy
from types import SimpleNamespace as NS
import pytest
from course_document import stable_hash
from course_exchange_lesson_ppt import lesson_ppt_projects
from course_exchange_ppt import public_ppt_options, freeze_published_ppts

@pytest.fixture
def lesson_export(monkeypatch, tmp_path):
    import teacher_content_projection, teacher_lesson_authoring, teaching_representations, course_revisions
    from storage import Storage
    manuscript={'manuscript_revision':'m1','quality_status':'passed','generation_contract_version':'script_ppt_bundle_v1'}
    state={'revision':'m1','status':'ready','source_state':'current','generation_contract_version':'script_ppt_bundle_v1', 'source_lesson_plan_revision_id':'p1','source_script_revision_id':'s1','source_material_revision':stable_hash([],prefix='pptrefs_'),'generated_representation_id':'r1','manuscript':manuscript}
    asset={'working_representation_id':'r1','source_state':'current','source_lesson_plan_revision_id':'p1','source_script_revision_id':'s1','ppt_manuscript_revision':'m1'}
    lesson={'working_revision_id':'p1','working_script_revision_id':'s1','revisions':[{'revision_id':'p1'}],'script_revisions':[{'revision_id':'s1'}],'ppt_manuscript':state,'ppt_assets':[asset]}
    binding=NS(material_evidence_ids=[],source_revisions={'block:b':'v1'})
    representation=NS(representation_id='r1',spec_id='sp1',status='ready',rebuild_required=False,source_bindings=[binding])
    content={'schema_version':'slide_deck_v6','pages':[], 'ppt_manuscript':deepcopy(manuscript)}
    spec=NS(spec_id='sp1',source_bindings=[binding],payload={'content':content})
    registry=NS(representations=[representation],specs=[spec])
    monkeypatch.setattr(teacher_content_projection,'authoring_repository',lambda storage:NS(load=lambda cid:{'lessons':{'lecture':lesson}}))
    monkeypatch.setattr(teacher_lesson_authoring,'teacher_lesson_v6_source',lambda *args,**kwargs:(NS(blocks=[NS(block_id='b',internal_revision='v1')]),{},'scope'))
    monkeypatch.setattr(course_revisions,'revision_vector_for_course',lambda *args:NS(revisions={'block:b':'v1'}))
    monkeypatch.setattr(teaching_representations,'teaching_representation_repository',NS(for_storage_scope=lambda *args:NS(load=lambda cid:registry)))
    course={'course_id':'c','authoring_surface':'teacher'}
    document={'sections':[{'section_id':'lecture','title':'Lecture'}], 'blocks':[{'block_id':'b','internal_revision':'v1','section_id':'lecture'}]}
    return Storage(str(tmp_path/'data')),course,document,lesson,representation,spec


def test_current_lecture_export_uses_existing_publication_and_private_part_filter(lesson_export, monkeypatch):
    from pptx import Presentation
    import slide_deck_v6_renderer
    store,course,doc,*_=lesson_export
    assert public_ppt_options(course,doc,storage=store)==[{'project_id':'lesson-ppt:lecture','title':'Lecture'}]
    calls=[]
    def export(snapshot,destination):
        calls.append(snapshot)
        p=Presentation();s=p.slides.add_slide(p.slide_layouts[6]);s.notes_slide.notes_text_frame.text='PRIVATE';p.save(destination)
    monkeypatch.setattr(slide_deck_v6_renderer,'export_slide_deck_v6_pptx',export)
    first=freeze_published_ppts(store,course,doc,['lesson-ppt:lecture'])
    assert first==freeze_published_ppts(store,course,doc,['lesson-ppt:lecture'])
    assert len(calls)==1 and first[0]['filename']=='Lecture.pptx'
    from zipfile import ZipFile
    assert all(b'PRIVATE' not in ZipFile(store._data_dir+'/published_ppts/'+first[0]['sha256']+'.pptx').read(n) for n in ZipFile(store._data_dir+'/published_ppts/'+first[0]['sha256']+'.pptx').namelist())

@pytest.mark.parametrize('change',['private_block','stale_block','plan','script','unbuilt','manuscript','private_upload','quality','representation','binding','material_binding','deck_revision'])
def test_incomplete_private_and_stale_lecture_decks_are_not_publication_options(lesson_export,change):
    store,course,doc,lesson,representation,spec=lesson_export
    state=lesson['ppt_manuscript']
    if change=='private_block':doc['blocks']=[]
    elif change=='stale_block':doc['blocks'][0]['internal_revision']='old'
    elif change=='plan':lesson['working_revision_id']='p2'
    elif change=='script':lesson['working_script_revision_id']='s2'
    elif change=='unbuilt':state['generated_representation_id']=''
    elif change=='manuscript':state['revision']='m2'
    elif change=='private_upload':state['source_material_revision']='private'
    elif change=='quality':state['manuscript']['quality_status']='blocked'
    elif change=='representation':representation.status='stale'
    elif change=='binding':spec.source_bindings[0].source_revisions={'block:b':'old'}
    elif change=='material_binding':spec.source_bindings[0].material_evidence_ids=['private-upload']
    else:spec.payload['content']['ppt_manuscript']['manuscript_revision']='old'
    assert lesson_ppt_projects(store,course,doc)=={}
    from fastapi import HTTPException
    with pytest.raises(HTTPException):freeze_published_ppts(store,course,doc,['lesson-ppt:lecture'])
