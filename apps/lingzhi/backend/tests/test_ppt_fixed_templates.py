"""Fixed fields -> saved manuscript -> editable PPTX, without model/network calls."""
import asyncio
from copy import deepcopy

import pytest
from pptx import Presentation

from course_document import CourseBlock, CourseDocument, CourseSection
from course_presentation_graph import compile_course_presentation_graph
from ppt_fixed_draft import lower_fixed_response
from ppt_fixed_templates import VERSION, compile_fixed_template
from ppt_teaching_content import PageTeachingV2
from ppt_teaching_manuscript import compile_teaching_manuscript, template_for_manuscript
from ppt_teaching_planner import normalize_page_response, plan_teaching_manuscript
from ppt_page_scene import resolve_page_scenes


TEXT = "课堂学习。先准备，再执行，最后检查。串行逐项执行，并行同时执行。相同任务条件下比较执行方式。为什么先准备？明确目标。公式为 x=1。"


def field(text):
    return {"text": text, "sources": [{"block_id": "b", "quote": TEXT}]}


def source():
    return {"b": {"block_id": "b", "block_revision": "r", "full_text": TEXT}}


def fixture():
    document = CourseDocument(course_id="fixed-test", title="课堂学习", document_revision="doc",
        sections=[CourseSection(section_id="s", title="课堂学习", position=0)],
        blocks=[CourseBlock(block_id="b", section_id="s", position=0, payload={"markdown": TEXT}, internal_revision="r")])
    graph = compile_course_presentation_graph(document, teaching_plan={})
    return document, graph, compile_fixed_template("qizhi-classroom")


def response(slug):
    base = {"title": "课堂学习", "notes": "详细讲解保存在教师备注。"}
    if slug in {"cover", "section"}:
        return {**base, "subtitle": field("明确目标")}
    if slug == "comparison":
        return {**base, "condition": field("相同任务条件"), "left_subject": field("串行"), "right_subject": field("并行"),
            "rows": [{"dimension": field("执行方式"), "left": field("逐项执行"), "right": field("同时执行")}],
            "conclusion": field("比较执行方式")}
    if slug == "flow":
        return {**base, "steps": [field(x) for x in ("准备", "执行", "检查")]}
    if slug == "question":
        return {**base, "question": field("为什么先准备？"), "answer": field("明确目标")}
    if slug == "formula":
        return {**base, "formula": {"sources": [{"block_id": "b", "quote": "x=1"}]}, "explanation": field("明确目标")}
    return {**base, "points": [field("明确目标"), field("先准备，再执行，最后检查") ]}


def lowered(slug):
    _, _, template = fixture()
    plan = {"layout_id": template.layout_id(slug), "page_goal": "明确目标"}
    return normalize_page_response(lower_fixed_response(response(slug), plan), source())


@pytest.mark.parametrize("slug", ["cover", "section", "agenda", "bullets", "summary", "comparison", "flow", "question", "formula"])
def test_every_fixed_form_resolves_to_auditable_editable_objects(slug, tmp_path):
    from ppt_native_scene import render_scene, audit_scene
    from pptx.util import Pt
    _, _, template = fixture()
    value = lowered(slug)
    scenes = resolve_page_scenes(page_id="p", title=value["title"], content=PageTeachingV2.model_validate(value["teaching"]),
        layout=template.get_layout(template.layout_id(slug)), template=template, source_document_revision="doc")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(960), Pt(540)
    for scene in scenes:
        render_scene(prs.slides.add_slide(prs.slide_layouts[6]), scene)
    path = tmp_path / "fixed.pptx"
    prs.save(path)
    readback = Presentation(path)
    for slide, scene in zip(readback.slides, scenes, strict=True):
        audit_scene(slide, scene)
    if slug == "comparison":
        cells = [o for o in scenes[0].objects if o.slot_id.startswith("cell.")]
        assert cells[0].y == cells[1].y and cells[0].x < cells[1].x
    if slug == "flow":
        assert len(scenes[0].edges) == 2
        nodes = [o for o in scenes[0].objects if o.element_id]
        assert len({o.y for o in nodes}) == 1
    if slug == "question":
        assert len(scenes) == 2
        assert "answer" not in {o.element_id for o in scenes[0].objects}
        assert "answer" in {o.element_id for o in scenes[1].objects}


def test_fixed_capacity_rejects_overflow_without_font_shrinking():
    _, _, template = fixture()
    value = lowered("bullets")
    value["teaching"]["elements"][0]["text"] = "明确目标" * 80
    with pytest.raises(ValueError, match="capacity"):
        resolve_page_scenes(page_id="p", title=value["title"], content=PageTeachingV2.model_validate(value["teaching"]),
            layout=template.get_layout(template.layout_id("bullets")), template=template, source_document_revision="doc")


def test_fixed_manuscript_roundtrip_edit_and_export(tmp_path):
    from ppt_teaching_manuscript import revise_teaching_manuscript
    from slide_deck_v6 import compile_slide_deck_v6_from_manuscript
    from slide_deck_v6_renderer import export_slide_deck_v6_pptx
    doc, graph, template = fixture()
    value = lowered("comparison")
    manuscript = compile_teaching_manuscript(doc, graph, template, {"central_question": "比较执行方式"}, [{
        **value, "page_id": "p", "teaching_unit_id": graph.units[0].teaching_unit_id, "source_block_ids": ["b"]}])
    assert template_for_manuscript(manuscript).template_digest == template.template_digest
    revised = revise_teaching_manuscript(manuscript, [{"page_id": "p", "title": "执行方式"}])
    assert revised.pages[0].resolved_scenes[0].objects[0].text == "执行方式"
    deck = compile_slide_deck_v6_from_manuscript(doc, graph, revised, template)
    path = export_slide_deck_v6_pptx(deck, tmp_path / "complete.pptx")
    assert TEXT in Presentation(path).slides[0].notes_slide.notes_text_frame.text


def test_fixed_planner_parallel_failure_recovery_and_selected_forms():
    from slide_deck_v6_models import V6BuildError
    doc, graph, template = fixture()
    saved, active, peak, requests = {}, 0, 0, []

    async def save(value, event):
        saved.clear()
        saved.update(deepcopy(value))

    async def planner(request):
        nonlocal active, peak
        if request["teaching_request"] == "narrative":
            assert not any(x["layout_id"].endswith("/figure") for x in request["layout_capabilities"]["available_layouts"])
            return {"narrative_brief": {"central_question": "明确目标"}, "pacing": {"max_physical_pages": 8, "rationale": "完整呈现教学任务"},
                "pages": [{"page_id": f"p{i}", "title": "课堂学习", "page_goal": goal, "teaching_unit_id": graph.units[0].teaching_unit_id,
                           "source_block_ids": ["b"], "layout_id": template.layout_id("bullets")} for i, goal in enumerate(("准备", "执行", "检查"))]}
        assert request["teaching_request"] == "fixed_fields"
        assert "composition_notes" not in str(request["response_contract"])
        requests.append(request)
        active += 1
        peak = max(peak, active)
        await asyncio.sleep(0.01)
        active -= 1
        if request["page"]["page_goal"] == "执行":
            raise RuntimeError("provider failed")
        return response("bullets")

    with pytest.raises(V6BuildError):
        asyncio.run(plan_teaching_manuscript(doc, graph, template, planner, on_checkpoint=save))
    assert peak == 3 and set(saved["pages"]) == {"p0", "p2"}
    retried = []
    async def recover(request):
        retried.append(request["page"]["page_goal"])
        return response("bullets")
    manuscript, _ = asyncio.run(plan_teaching_manuscript(doc, graph, template, recover, checkpoint=saved))
    assert retried == ["执行"]
    assert [p.page_id for p in manuscript.pages] == ["p0", "p1", "p2"]
    assert manuscript.template_version == VERSION


def test_new_default_template_and_historical_template_lock(monkeypatch):
    from routers import teacher_lesson_authoring as routes
    from template_layout_contract import compile_builtin_template_layout_contract_v1
    monkeypatch.delenv("PPT_THREE_STAGE_ENABLED", raising=False)
    template = routes._resolve_teacher_v6_template(routes.TeacherLessonV6BuildRequest(), "teacher")
    assert template.template_version == VERSION
    for candidate in (template, compile_builtin_template_layout_contract_v1("academic-editorial")):
        locked = routes._resolve_locked_teacher_v6_template({"theme": candidate.theme_id,
            "template_version": candidate.template_version, "template_digest": candidate.template_digest}, "teacher")
        assert locked.template_digest == candidate.template_digest


def test_fixed_form_repair_and_confirmed_build_use_existing_orchestrator(tmp_path):
    from slide_deck_v6_models import PptManuscriptV1
    from slide_deck_v6_orchestrator import SlideDeckV6Orchestrator, SlideDeckV6CandidateRepository
    from teaching_representations import TeachingRepresentationRepository
    doc, graph, template = fixture()
    calls = []
    async def planner(request):
        calls.append(request["teaching_request"])
        if request["teaching_request"] == "narrative":
            return {"narrative_brief": {"central_question": "明确目标"}, "pacing": {"max_physical_pages": 2, "rationale": "呈现一个教学任务"},
                "pages": [{"page_id": "p", "title": "课堂学习", "page_goal": "明确目标", "teaching_unit_id": graph.units[0].teaching_unit_id,
                           "source_block_ids": ["b"], "layout_id": template.layout_id("bullets")}]}
        if calls.count("fixed_fields") == 1:
            return {"title": "课堂学习", "points": [], "notes": "补充讲解"}
        assert request["validation_error"]
        value = response("bullets")
        quote_id = request["literal_source_ranges"][0]["quote_id"]
        for point in value["points"]:
            point["sources"] = [{"quote_id": quote_id}]
        return value
    async def forbidden(request):
        pytest.fail("confirmed export called an AI planner")
    orchestrator = SlideDeckV6Orchestrator(representation_repository=TeachingRepresentationRepository(tmp_path / "representations"),
        candidate_repository=SlideDeckV6CandidateRepository(tmp_path / "candidates"), progress_root=tmp_path / "progress")
    args = dict(document=doc, course_data={}, mode="teaching", theme=template.theme_id,
        visual_planner=forbidden, source_revision_provider=lambda: doc.document_revision,
        template_contract=template, publish_result=False)
    draft = asyncio.run(orchestrator.build(task_id="draft", story_planner=planner, manuscript_only=True, **args))
    assert draft["status"] == "manuscript_ready" and calls == ["narrative", "fixed_fields", "fixed_fields"]
    manuscript = PptManuscriptV1.model_validate(draft["ppt_manuscript"])
    final = asyncio.run(orchestrator.build(task_id="final", story_planner=forbidden, confirmed_manuscript=manuscript, **args))
    assert final["status"] == "v6_ready"


def test_fixed_figure_uses_the_existing_adopted_asset_catalog():
    _, _, template = fixture()
    value = lower_fixed_response({"title": "课堂学习", "notes": "观察图片", "asset_id": "asset",
        "caption": field("明确目标")}, {"layout_id": template.layout_id("figure"), "page_goal": "明确目标"},
        [{"kind": "image", "assets": [{"asset_id": "asset", "sha256": "digest"}]}])
    image = value["elements"][0]
    assert image["kind"] == "image" and image["asset_id"] == "asset" and image["asset_digest"] == "digest"
    with pytest.raises(ValueError, match="asset_unavailable"):
        lower_fixed_response({"title": "课堂学习", "notes": "观察图片", "asset_id": "unknown", "caption": field("明确目标")},
            {"layout_id": template.layout_id("figure"), "page_goal": "明确目标"})


def test_authored_styles_survive_export_and_tampering_is_detected(tmp_path):
    from ppt_native_scene import render_scene, audit_scene
    from pptx.util import Pt
    from pptx.dml.color import RGBColor
    _, _, template = fixture()
    value = lowered("flow")
    scene = resolve_page_scenes(page_id="p", title=value["title"], content=PageTeachingV2.model_validate(value["teaching"]),
        layout=template.get_layout(template.layout_id("flow")), template=template, source_document_revision="doc")[0]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(960), Pt(540)
    render_scene(prs.slides.add_slide(prs.slide_layouts[6]), scene)
    path = tmp_path / "centered.pptx"
    prs.save(path)
    slide = Presentation(path).slides[0]
    audit_scene(slide, scene)
    node = next(s for s in slide.shapes if s.name == "teaching:item-0")
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    assert node.text_frame.vertical_anchor == MSO_ANCHOR.MIDDLE
    assert node.text_frame.paragraphs[0].alignment == PP_ALIGN.CENTER
    node.text_frame.paragraphs[0].font.color.rgb = RGBColor.from_string("FFFFFF")
    with pytest.raises(ValueError, match="text_style_mismatch"):
        audit_scene(slide, scene)


def test_v1_template_and_scene_digest_remain_readable():
    from course_document import stable_hash
    from ppt_fixed_templates import LEGACY_VERSION
    from ppt_page_scene import ResolvedPageScene, verify_scene
    _, _, current = fixture()
    legacy = compile_fixed_template(current.theme_id, version=LEGACY_VERSION)
    assert legacy.template_digest == "tmpl_bcf315cab5f4ec471efa3399"
    assert len(legacy.layouts) == 10 and len(current.layouts) == 18
    data = normalize_page_response(lower_fixed_response(response("flow"),
        {"layout_id": legacy.layout_id("flow"), "page_goal": "明确目标"}), source())
    scene = resolve_page_scenes(page_id="p", title=data["title"], content=PageTeachingV2.model_validate(data["teaching"]),
        layout=legacy.get_layout(legacy.layout_id("flow")), template=legacy, source_document_revision="doc")[0]
    stored = scene.model_dump(mode="json")
    assert all("text_align" not in o and "vertical_align" not in o and "visible_with" not in o for o in stored["objects"])
    assert stable_hash({k: v for k, v in stored.items() if k != "scene_digest"}, prefix="scene_") == stored["scene_digest"]
    verify_scene(ResolvedPageScene.model_validate(stored))


def test_explanation_headings_keep_sources_and_fit_their_own_fields():
    _, _, template = fixture()
    data = response("bullets")
    data["points"][0]["heading"] = "准备"
    value = normalize_page_response(lower_fixed_response(data,
        {"layout_id": template.layout_id("bullets"), "page_goal": "明确目标"}), source())
    content = PageTeachingV2.model_validate(value["teaching"])
    scene = resolve_page_scenes(page_id="p", title=value["title"], content=content,
        layout=template.get_layout(template.layout_id("bullets")), template=template, source_document_revision="doc")[0]
    heading = next(o for o in scene.objects if o.element_id == "item-0-heading")
    body = next(o for o in scene.objects if o.element_id == "item-0")
    assert heading.font_size > body.font_size and heading.y + heading.height <= body.y
    assert {s.block_id for e in content.elements for s in e.sources} == {"b"}


@pytest.mark.parametrize("slug", ["chart", "code"])
def test_data_and_code_remain_source_exact_in_real_pptx(slug, tmp_path):
    from ppt_native_scene import render_scene, audit_scene
    from pptx.util import Pt
    _, _, template = fixture()
    code = "for item in items:\n    print(item)"
    raw = {"b": {"block_id": "b", "block_revision": "r", "full_text": "甲 10，乙 20，单位：人。\n" + code}}
    exact = lambda quote: {"sources": [{"block_id": "b", "quote": quote}]}
    text = lambda value: {"text": value, **exact(raw["b"]["full_text"])}
    data = {"title": "课堂学习", "notes": "使用示例数据"}
    data.update({"unit": exact("人"), "points": [{"label": text("甲"), "value": exact("10")}, {"label": text("乙"), "value": exact("20")}]}
                if slug == "chart" else {"code": exact(code), "explanation": text("逐项执行")})
    value = normalize_page_response(lower_fixed_response(data,
        {"layout_id": template.layout_id(slug), "page_goal": "明确目标"}), raw)
    scene = resolve_page_scenes(page_id="p", title=value["title"], content=PageTeachingV2.model_validate(value["teaching"]),
        layout=template.get_layout(template.layout_id(slug)), template=template, source_document_revision="doc")[0]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(960), Pt(540)
    render_scene(prs.slides.add_slide(prs.slide_layouts[6]), scene)
    path = tmp_path / f"{slug}.pptx"
    prs.save(path)
    slide = Presentation(path).slides[0]
    audit_scene(slide, scene)
    if slug == "chart":
        bars = [s for s in slide.shapes if s.name.startswith("teaching:chart-bar-")]
        assert bars[0].left == bars[1].left and bars[1].width == bars[0].width * 2
        assert {e.text for e in PageTeachingV2.model_validate(value["teaching"]).elements if e.kind == "data"} == {"10", "20"}
    else:
        assert next(s.text for s in slide.shapes if s.name == "teaching:code") == code


def test_fixed_request_preserves_adopted_materials_and_lesson_context():
    from ppt_fixed_draft import invoke_fixed_form
    _, _, template = fixture()
    async def planner(request):
        import json
        json.dumps(request)
        assert request["page"]["page_id"] == "page-3"
        assert request["lesson_context"]["central_question"] == "明确目标"
        assert request["accepted_question_bank_items"] == [{"question": "为什么先准备？"}]
        assert request["page_sequence"] == [{"title": "课堂学习", "page_goal": "明确目标"}]
        return response("bullets")
    result = asyncio.run(invoke_fixed_form(planner, {"page": {"page_id": "page-3", "title": "课堂学习", "page_goal": "明确目标", "layout_id": template.layout_id("bullets")},
        "narrative_brief": {"central_question": "明确目标"}, "accepted_question_bank_items": [{"question": "为什么先准备？"}],
        "page_sequence": [{"title": "课堂学习", "page_goal": "明确目标"}]}))
    assert result["expression_kind"] == "evidence"


def test_authored_cover_has_no_hairline_and_figure_keeps_adopted_image(tmp_path):
    from ppt_native_scene import render_scene, audit_scene
    from .test_ppt_adopted_visuals import adopted_fixture
    from ppt_adopted_visuals import bind_adopted_assets, current_visual_catalog
    from pptx.util import Pt
    _, _, template = fixture()
    value = lowered("cover")
    scene = resolve_page_scenes(page_id="p", title=value["title"], content=PageTeachingV2.model_validate(value["teaching"]),
        layout=template.get_layout(template.layout_id("cover")), template=template, source_document_revision="doc")[0]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(960), Pt(540)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    render_scene(slide, scene)
    assert scene.background == "2B5876"
    title = next(s for s in slide.shapes if s.name == "teaching:title")
    assert title._element.xpath('./p:spPr/a:ln/a:noFill')
    assert title.text_frame.paragraphs[0].font.size == Pt(44)
    assets, asset, item, _, raw = adopted_fixture(tmp_path)
    catalog = current_visual_catalog([item], course_id='course', script_revision_id='script-r1', sources=raw, asset_repository=assets)
    data = {"title": "执行方式", "notes": "观察图示", "asset_id": asset.asset_id,
        "caption": {"text": "执行方式示意", "sources": [{"block_id": "b", "quote": raw["b"]}]}}
    value = normalize_page_response(lower_fixed_response(data,
        {"layout_id": template.layout_id("figure"), "page_goal": "理解执行方式"}, catalog),
        {"b": {"block_id": "b", "block_revision": "r", "full_text": raw["b"]}})
    content = bind_adopted_assets(PageTeachingV2.model_validate(value["teaching"]), catalog, {"b"})
    scene = resolve_page_scenes(page_id="figure", title=value["title"], content=content,
        layout=template.get_layout(template.layout_id("figure")), template=template, source_document_revision="doc")[0]
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    render_scene(slide, scene, assets=assets)
    audit_scene(slide, scene)
    assert next(s for s in slide.shapes if s.name == "teaching:image").image.blob == assets.resolve(asset.asset_id).read_bytes()
