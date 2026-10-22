"""Real HTTP between two isolated API datasets; no shared user/course database."""

from contextvars import ContextVar
from copy import deepcopy
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
import threading
from pathlib import Path

from fastapi import FastAPI
from fastapi.testclient import TestClient
import pytest
import requests
from starlette.middleware.base import BaseHTTPMiddleware

from course_access import CourseOwnershipMiddleware
from course_document import CourseBlock, CourseDocument, CourseSection, refresh_document_revision
from course_exchange import ExchangeConfig, ExchangeRepository, INTERNAL_PREFIX
from qizhi_auth import QizhiIdentityMiddleware, QizhiIdentityVerifier
from storage import Storage
from study.auth import LocalStudyAuth
from study.identity_middleware import LocalStudyIdentityMiddleware

node_context = ContextVar("exchange_test_node", default="teacher")


@pytest.fixture
def pair(tmp_path, monkeypatch):
    monkeypatch.setenv("LINGZHI_LOCAL_STUDY_AUTH", "true")
    monkeypatch.delenv("QIZHI_AUTH_VERIFY_URL", raising=False)
    stores = {name: Storage(str(tmp_path / name)) for name in ("teacher", "student")}
    configs = {}

    class NodeStorage:
        def __getattr__(self, name):
            return getattr(stores[node_context.get()], name)

    proxy = NodeStorage()
    monkeypatch.setattr(ExchangeConfig, "environment", classmethod(lambda cls: configs[node_context.get()]))
    from routers import (
        course_exchange as api,
        course_learning_access as grants,
        courses,
        practice,
        learning_assets,
        study,
    )
    import dependencies, learning_events, learning_progress, learning_continuation

    for module in (api, grants, courses, practice, study, dependencies, learning_events):
        monkeypatch.setattr(module, "storage", proxy)
    from practice_attempts import PracticeAttemptRepository
    import practice_attempts

    attempts = PracticeAttemptRepository(tmp_path / "student" / "practice_attempts")
    for module in (practice, practice_attempts, learning_progress, learning_continuation):
        monkeypatch.setattr(module, "practice_attempt_repository", attempts)
    monkeypatch.setattr(practice, "_active_question_bank", lambda *a: None)
    monkeypatch.setattr(practice.question_bank_rebuild_job_repository, "latest_for_course", lambda *a: None)
    monkeypatch.setattr(practice.learning_asset_repository, "load_bundle", lambda *a: None)
    monkeypatch.setattr(practice, "advance_workflow_after_grade", lambda *a, **kw: {})
    monkeypatch.setattr(practice, "workflow_view", lambda *a, **kw: {})
    monkeypatch.setattr(learning_continuation, "workflow_view", lambda *a, **kw: {})
    monkeypatch.setattr(learning_continuation.learning_record_repository, "list", lambda *a: [])
    monkeypatch.setattr(learning_continuation.learning_snapshot_repository, "load", lambda *a: {})

    class UnavailableDiagnosis:
        async def diagnose_answer(self, *a):
            raise RuntimeError("No model in deterministic HTTP acceptance")

    monkeypatch.setattr(practice, "practice_analysis_service", UnavailableDiagnosis())
    doc = refresh_document_revision(
        CourseDocument(
            course_id="source-course",
            title="Arithmetic",
            sections=[CourseSection(section_id="lecture", title="Addition", position=0, level=1)],
            blocks=[
                CourseBlock(
                    block_id="public",
                    section_id="lecture",
                    position=0,
                    payload={"markdown": "2 + 3 = 5.", "private_note": "SECRET"},
                ),
                CourseBlock(
                    block_id="private",
                    section_id="lecture",
                    position=1,
                    payload={"markdown": "SECRET"},
                    visibility_rule={"audience": "teacher"},
                ),
            ],
        )
    )
    course = {
        "course_id": "source-course",
        "owner_id": "teacher",
        "authoring_surface": "generated",
        "course_schema_version": "course_document_v1",
        "course_document": doc.model_dump(mode="json"),
        "course_name": "Arithmetic",
        "current_course_version_id": doc.document_revision,
        "learning_assets": {
            "questions": [
                {
                    "asset_id": "question",
                    "revision_id": "q1",
                    "node_id": "lecture",
                    "question_type": "single_choice",
                    "practice_level": "mastery_check",
                    "prompt": "What is 2 + 3?",
                    "options": [{"id": "A", "text": "5"}, {"id": "B", "text": "6"}],
                    "input_contract": {"mode": "choice", "required": True, "selection": {"multiple": False}},
                    "answer_spec": {"type": "choice", "correct_option_id": "A", "pass_score": 70},
                    "grading_policy": {"method": "choice"},
                    "hint_contract": {"levels": [{"level": 1, "kind": "strategy", "content": "Add the two numbers."}]},
                }
            ],
            "mastery_criteria": [],
            "checklist": [],
        },
    }
    stores["teacher"].save_course_sync("source-course", course)
    account = LocalStudyAuth(stores["student"]).login("student_a", "Testing2026!", register=True)
    other = LocalStudyAuth(stores["student"]).login("student_b", "Testing2026!", register=True)
    clients = {}
    servers, threads = [], []
    for name in ("teacher", "student"):
        app = FastAPI()
        for router in (
            api.router,
            grants.router,
            courses.router,
            practice.router,
            learning_assets.router,
            study.router,
        ):
            app.include_router(router, prefix="/api")
        app.add_middleware(CourseOwnershipMiddleware, course_storage=proxy)
        app.add_middleware(LocalStudyIdentityMiddleware, course_storage=proxy)
        app.add_middleware(QizhiIdentityMiddleware, verifier=QizhiIdentityVerifier(""))

        class NodeContext(BaseHTTPMiddleware):
            def __init__(self, app, node):
                super().__init__(app)
                self.node = node

            async def dispatch(self, request, call_next):
                token = node_context.set(self.node)
                try:
                    return await call_next(request)
                finally:
                    node_context.reset(token)

        app.add_middleware(NodeContext, node=name)
        clients[name] = TestClient(app)

        def handler_for(client):
            class Handler(BaseHTTPRequestHandler):
                def do_POST(self):
                    body = self.rfile.read(int(self.headers.get("Content-Length", "0")))
                    result = client.post(self.path, content=body, headers=dict(self.headers))
                    self.send_response(result.status_code)
                    for key in ("Content-Type", "X-Lingzhi-Node"):
                        if key in result.headers:
                            self.send_header(key, result.headers[key])
                    self.end_headers()
                    self.wfile.write(result.content)

                def log_message(self, *args):
                    pass

            return Handler

        server = ThreadingHTTPServer(("127.0.0.1", 0), handler_for(clients[name]))
        servers.append(server)
    token = "only-a-synthetic-test-secret-" + "x" * 32
    configs["teacher"] = ExchangeConfig(
        "teacher", "teacher-test", "student-test", f"http://127.0.0.1:{servers[1].server_port}", token
    )
    configs["student"] = ExchangeConfig(
        "student", "student-test", "teacher-test", f"http://127.0.0.1:{servers[0].server_port}", token
    )
    for server in servers:
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        threads.append(thread)

    class Pair:
        def request(self, node, method, path, **kwargs):
            return clients[node].request(method, path, **kwargs)

        def repository(self, node):
            return ExchangeRepository(stores[node], configs[node])

        def publish(self, **extra):
            view = self.request(
                "teacher", "GET", "/api/teacher/courses/source-course/exchange", headers={"X-User-Id": "teacher"}
            )
            assert view.status_code == 200, view.text
            response = self.request(
                "teacher",
                "POST",
                "/api/teacher/courses/source-course/exchange/publish",
                headers={"X-User-Id": "teacher"},
                json={"expected_revision": view.json()["current_revision"], **extra},
            )
            assert response.status_code == 200, response.text
            return response.json()

    result = Pair()
    result.stores = stores
    result.configs = configs
    result.account = account
    result.other = other
    result.actor = account["user"]["id"]
    result.headers = {"Authorization": "Bearer " + account["token"]}
    result.attempts = attempts
    result.course = course
    yield result
    for server in servers:
        server.shutdown()
        server.server_close()
    for thread in threads:
        thread.join(timeout=3)
    for client in clients.values():
        client.close()


def enroll(pair):
    granted = pair.request(
        "teacher",
        "POST",
        "/api/teacher/courses/source-course/learners",
        headers={"X-User-Id": "teacher"},
        json={"local_username": "student_a"},
    )
    assert granted.status_code == 200, granted.text
    assert granted.json()["learner_id"] == pair.actor
    pair.publish()
    response = pair.request("student", "GET", "/api/study/courses", headers=pair.headers)
    assert response.status_code == 200, response.text
    return response.json()[0]


def test_separate_data_http_publish_read_and_revocation(pair):
    from question_public_guard import assert_no_answer_leak

    course = enroll(pair)
    assert not pair.stores["teacher"].load_data("study_accounts.json")
    assert not pair.stores["student"].load_course(course["course_id"])
    assert course["question_count"] == 1
    read = pair.request("student", "GET", f"/api/courses/{course['course_id']}/document", headers=pair.headers)
    assert read.status_code == 200, read.text
    assert "SECRET" not in read.text and "2 + 3 = 5" in read.text
    questions = pair.request(
        "student", "GET", f"/api/courses/{course['course_id']}/practice?node_id=lecture", headers=pair.headers
    )
    assert questions.status_code == 200, questions.text
    assert_no_answer_leak(questions.json())
    unauthorized = pair.request(
        "student",
        "GET",
        f"/api/courses/{course['course_id']}/document",
        headers={"Authorization": "Bearer " + pair.other["token"]},
    )
    assert unauthorized.status_code == 404
    assert (
        pair.request(
            "student", "POST", f"/api/courses/{course['course_id']}/lesson-authoring", headers=pair.headers
        ).status_code
        == 403
    )
    pair.request(
        "teacher",
        "DELETE",
        "/api/teacher/courses/source-course/learners/" + pair.actor,
        headers={"X-User-Id": "teacher"},
    )
    assert (
        pair.request("student", "GET", f"/api/courses/{course['course_id']}/document", headers=pair.headers).status_code
        == 404
    )
    assert pair.repository("student").references(pair.actor)


def test_fixed_publication_and_explicit_student_update(pair):
    course = enroll(pair)
    original = pair.repository("student").read(pair.actor, course["course_id"])
    raw = deepcopy(pair.course)
    raw["course_document"]["blocks"][0]["payload"]["markdown"] = "Updated public content."
    raw["course_document"] = refresh_document_revision(
        CourseDocument.model_validate(raw["course_document"])
    ).model_dump(mode="json")
    pair.stores["teacher"].save_course_sync("source-course", raw)
    # Teacher drafting does not rewrite an already published release.
    assert pair.repository("student").read(pair.actor, course["course_id"]) == original
    newer = pair.publish()
    listing = pair.request("student", "GET", "/api/study/courses", headers=pair.headers).json()[0]
    assert listing["update_available"] and listing["revision_id"] == course["revision_id"]
    updated = pair.request(
        "student",
        "POST",
        f"/api/study/courses/{course['course_id']}/exchange/update",
        headers=pair.headers,
        json={"expected_release_id": course["revision_id"]},
    )
    assert updated.status_code == 200, updated.text
    assert updated.json()["release_id"] == newer["release_id"]
    assert (
        pair.repository("student").read(pair.actor, course["course_id"], release_id=course["revision_id"]) == original
    )


def test_course_feedback_reply_and_idempotent_delivery(pair):
    course = enroll(pair)
    url = f"/api/study/courses/{course['course_id']}/exchange/feedback"
    request = {
        "request_id": "question-1",
        "expected_release_id": course["revision_id"],
        "text": "Could you explain addition again?",
    }
    first = pair.request("student", "POST", url, headers=pair.headers, json=request)
    assert first.status_code == 200 and first.json()["status"] == "delivered", first.text
    assert pair.request("student", "POST", url, headers=pair.headers, json=request).json() == first.json()
    teacher = pair.request(
        "teacher", "GET", "/api/teacher/courses/source-course/exchange/feedback", headers={"X-User-Id": "teacher"}
    ).json()
    assert teacher["count"] == 1
    item = teacher["items"][0]
    reply = pair.request(
        "teacher",
        "POST",
        "/api/teacher/courses/source-course/exchange/feedback/" + item["feedback_id"] + "/reply",
        headers={"X-User-Id": "teacher"},
        json={"expected_revision": item["revision"], "text": "Count two objects, then three more."},
    )
    assert reply.status_code == 200
    history = pair.request("student", "GET", url, headers=pair.headers).json()
    assert history["items"][0]["reply"] == "Count two objects, then three more."
    assert (
        pair.request(
            "teacher", "GET", "/api/teacher/courses/source-course/exchange/feedback", headers={"X-User-Id": "foreign"}
        ).status_code
        == 404
    )


def test_remote_grading_wrong_retry_and_teacher_receipt(pair):
    course = enroll(pair)
    base = f"/api/courses/{course['course_id']}/practice"
    created = pair.request("student", "POST", base + "/attempts", headers=pair.headers, json={"task_revision_id": "q1"})
    assert created.status_code == 200, created.text
    original = created.json()["attempt"]

    def submit(attempt, selected, request_id):
        result = pair.request(
            "student",
            "POST",
            base + "/attempts/" + attempt["attempt_id"] + "/submit",
            headers=pair.headers,
            json={
                "expected_revision": attempt["revision"],
                "answer_payload": {"selected_option_id": selected},
                "active_seconds": 9,
                "request_id": request_id,
            },
        )
        assert result.status_code == 200, result.text
        return result.json()

    wrong = submit(original, "B", "submit-1")
    assert wrong["result"]["passed"] is False and wrong["feedback_delivery"]["status"] == "delivered", wrong
    retry = pair.request(
        "student",
        "POST",
        base + "/attempts",
        headers=pair.headers,
        json={
            "task_revision_id": "q1",
            "resume": False,
            "practice_intent": "targeted_retry",
            "origin_attempt_id": original["attempt_id"],
        },
    )
    assert retry.status_code == 200, retry.text
    correct = submit(retry.json()["attempt"], "A", "submit-2")
    assert correct["result"]["passed"] is True and correct["result"]["score"] == 100, correct
    assert (
        pair.request("student", "GET", base + "/history?view=needs_review", headers=pair.headers).json()["attempts"]
        == []
    )
    assert len(pair.attempts.list(pair.actor, course["course_id"])) == 2
    assert not list((Path(pair.stores["teacher"]._data_dir) / "practice_attempts").glob("**/*.json"))
    feedback = pair.request(
        "teacher", "GET", "/api/teacher/courses/source-course/exchange/feedback", headers={"X-User-Id": "teacher"}
    ).json()
    assert feedback["count"] == 2
    assert {i["assessment"]["result"]["passed"] for i in feedback["items"]} == {True, False}
    assert all(i["assessment"]["task"]["prompt"] == "What is 2 + 3?" for i in feedback["items"])


def test_peer_authentication_is_not_a_browser_student_session(pair):
    response = requests.post(
        pair.configs["teacher"].peer_url + INTERNAL_PREFIX + "identity",
        json={"username": "student_a"},
        headers={"Authorization": "Bearer " + pair.account["token"], "X-Lingzhi-Node": "teacher-test"},
        timeout=5,
    )
    assert response.status_code == 401


def test_owner_transfer_does_not_transfer_private_feedback(pair):
    course = enroll(pair)
    response = pair.request(
        "student",
        "POST",
        f"/api/study/courses/{course['course_id']}/exchange/feedback",
        headers=pair.headers,
        json={
            "request_id": "old-owner-question",
            "expected_release_id": course["revision_id"],
            "text": "Question for original teacher",
        },
    )
    assert response.status_code == 200
    raw = deepcopy(pair.course)
    raw["owner_id"] = "new-teacher"
    pair.stores["teacher"].save_course_sync("source-course", raw)
    visible = pair.request(
        "teacher", "GET", "/api/teacher/courses/source-course/exchange/feedback", headers={"X-User-Id": "new-teacher"}
    )
    assert visible.status_code == 200 and visible.json()["count"] == 0
    state = pair.request(
        "teacher", "GET", "/api/teacher/courses/source-course/exchange", headers={"X-User-Id": "new-teacher"}
    )
    assert state.json()["status"] == "unpublished"
    reply = pair.request(
        "teacher",
        "POST",
        "/api/teacher/courses/source-course/exchange/feedback/" + response.json()["feedback_id"] + "/reply",
        headers={"X-User-Id": "new-teacher"},
        json={"expected_revision": 1, "text": "Should not see this"},
    )
    assert reply.status_code == 404
    assert (
        pair.request("student", "GET", f"/api/courses/{course['course_id']}/document", headers=pair.headers).status_code
        == 404
    )


def test_published_knowledge_keeps_public_bindings_and_excludes_private_sources(pair):
    from course_exchange import public_knowledge_tree

    raw = deepcopy(pair.course)
    raw["course_knowledge_base"] = {
        "schema_version": "course_knowledge_base_v2",
        "lifecycle_status": "active",
        "revision_id": "tree-a",
        "knowledge_points": [
            {
                "knowledge_id": "addition",
                "name": "Addition",
                "statement": "Combine quantities",
                "revision_id": "k1",
                "private_note": "SECRET",
            }
        ],
        "relations": [
            {
                "relation_id": "r1",
                "source_knowledge_id": "addition",
                "target_knowledge_id": "count",
                "relation_type": "prerequisite",
                "reason": "Counting supports addition",
                "source_refs": ["SECRET"],
            }
        ],
        "section_bindings": {
            "lecture": {"knowledge_node_ids": ["addition"], "capability_point_ids": [], "private_note": "SECRET"}
        },
    }
    public = public_knowledge_tree(raw, raw["course_document"])
    assert public["section_bindings"]["lecture"]["knowledge_node_ids"] == ["addition"]
    assert public["relations"][0]["reason"] == "Counting supports addition"
    assert "SECRET" not in str(public)


def generic_handout_course(course):
    raw = deepcopy(course)
    raw["teacher_handouts"] = {"lecture": {"revision_id": "script-a"}}
    group = {"id": "arithmetic", "concept_group_id": "arithmetic", "parent_id": None, "name": "Arithmetic"}
    points = []
    for kid, block in (("addition", "public"), ("count", "public"), ("private-atom", "private")):
        points.append({
            "knowledge_id": kid, "revision_id": kid+"-v1", "name": kid,
            "statement": "Public concept" if block == "public" else "SECRET",
            "parent_id": "arithmetic", "section_refs": ["lecture"], "topic": "lecture",
            "granularity_status": "candidate", "private_note": "SECRET", "meta_answer": "SECRET",
            "source_refs": ["source-"+kid],
            "source_quotes": [{"evidence_id": "source-"+kid, "quote": "2 + 3 = 5.",
                               "handout_ref": {"lesson_id": "lecture", "script_revision_id": "script-a", "block_id": block}}],
        })
    raw["course_knowledge_base"] = {
        "schema_version": "course_knowledge_base_v2", "tree_schema": "knowledge_first_generic_v1",
        "status": "accepted", "revision_id": "tree-a", "source_mode": "handouts",
        "categories": [group], "concept_groups": [group], "knowledge_points": points,
        "relations": [
            {"id": "count-before-add", "source_id": "count", "target_id": "addition", "type": "prerequisite", "reason": "Counting supports addition", "private_note": "SECRET"},
            {"id": "private-edge", "source_id": "addition", "target_id": "private-atom", "type": "applies_to"},
        ],
        "lessons": [{"topic": "lecture", "title": "Arithmetic", "knowledge_ids": [p["knowledge_id"] for p in points]}],
        "source_handouts": [{"lesson_id": "lecture", "script_revision_id": "script-a"}],
        "sources": [{"evidence_id": "source-"+p["knowledge_id"]} for p in points],
        "bindings": [], "skill_units": [], "misconceptions": [], "mastery_criteria": [],
    }
    return raw


def test_generic_handout_tree_reaches_student_as_a_frozen_public_library(pair):
    from question_public_guard import assert_no_answer_leak

    raw = generic_handout_course(pair.course)
    pair.stores["teacher"].save_course_sync("source-course", raw)
    course = enroll(pair)
    path = f"/api/courses/{course['course_id']}/learning-assets"
    first = pair.request("student", "GET", path, headers=pair.headers)
    assert first.status_code == 200, first.text
    library = first.json()["assets"]["knowledge_library"][0]
    assert library["schema_version"] == "knowledge_library_view_v3"
    assert library["binding_revision_id"] == "tree-a"
    assert library["subject_id"] == course["course_id"]
    points = {n["knowledge_id"]: n for n in library["nodes"] if n["node_type"] == "knowledge_point"}
    assert set(points) == {"addition", "count"}
    assert all(n["covered_by_course"] and n["granularity_status"] == "candidate" for n in points.values())
    assert points["addition"]["parent_id"] == "arithmetic"
    edge = library["relations"][0]
    assert (edge["relation_id"], edge["source_knowledge_id"], edge["target_knowledge_id"], edge["relation_type"]) == (
        "count-before-add", "count", "addition", "prerequisite")
    assert "SECRET" not in first.text and "private-atom" not in first.text
    assert_no_answer_leak(first.json()["assets"])
    original = pair.repository("student").read(pair.actor, course["course_id"])

    raw["course_knowledge_base"]["revision_id"] = "tree-b"
    raw["course_knowledge_base"]["knowledge_points"][0]["statement"] = "Revised public concept"
    pair.stores["teacher"].save_course_sync("source-course", raw)
    assert pair.request("student", "GET", path, headers=pair.headers).json()["assets"]["knowledge_library"][0] == library
    newer = pair.publish()
    assert newer["release_id"] != course["revision_id"]
    assert pair.request("student", "GET", path, headers=pair.headers).json()["assets"]["knowledge_library"][0] == library
    updated = pair.request("student", "POST", f"/api/study/courses/{course['course_id']}/exchange/update",
                           headers=pair.headers, json={"expected_release_id": course["revision_id"]})
    assert updated.status_code == 200, updated.text
    current = pair.request("student", "GET", path, headers=pair.headers).json()["assets"]["knowledge_library"][0]
    assert current["binding_revision_id"] == "tree-b"
    assert next(n for n in current["nodes"] if n["knowledge_id"] == "addition")["statement"] == "Revised public concept"
    assert pair.repository("student").read(pair.actor, course["course_id"], release_id=course["revision_id"]) == original


@pytest.mark.parametrize("source_change", ["candidate", "handout_changed"])
def test_generic_public_knowledge_excludes_unadopted_or_changed_handout_points(pair, source_change):
    from course_exchange import public_knowledge_tree
    from course_learning_projection import public_document

    raw = generic_handout_course(pair.course)
    if source_change == "candidate":
        raw["course_knowledge_base"].update(status="candidate", lifecycle_status="active")
    else:
        raw["teacher_handouts"]["lecture"]["revision_id"] = "script-b"
    public = public_knowledge_tree(raw, public_document(raw))
    assert not public.get("knowledge_points")
    assert "private-atom" not in str(public) and "SECRET" not in str(public)


def test_pending_feedback_survives_outage_and_retry_is_exactly_once(pair, monkeypatch):
    from course_exchange import ExchangeClient
    from fastapi import HTTPException

    course = enroll(pair)
    url = f"/api/study/courses/{course['course_id']}/exchange/feedback"
    original = ExchangeClient.call

    def offline(self, *args, **kwargs):
        raise HTTPException(503, "Synthetic network outage")

    monkeypatch.setattr(ExchangeClient, "call", offline)
    payload = {
        "request_id": "offline-question",
        "expected_release_id": course["revision_id"],
        "text": "Explain the example.",
    }
    saved = pair.request("student", "POST", url, headers=pair.headers, json=payload)
    assert saved.status_code == 200 and saved.json()["status"] == "pending", saved.text
    history = pair.request("student", "GET", url, headers=pair.headers)
    assert history.status_code == 200 and not history.json()["replies_available"]
    assert history.json()["items"][0]["payload"]["text"] == payload["text"]
    assert (
        pair.request("student", "GET", f"/api/courses/{course['course_id']}/document", headers=pair.headers).status_code
        == 503
    )
    monkeypatch.setattr(ExchangeClient, "call", original)
    retry = url + "/" + saved.json()["feedback_id"] + "/retry"
    assert pair.request("student", "POST", retry, headers=pair.headers).json()["status"] == "delivered"
    assert pair.request("student", "POST", retry, headers=pair.headers).json()["status"] == "delivered"
    assert len(pair.repository("teacher").ledger()["received"]) == 1
    assert (
        pair.request("student", "POST", retry, headers={"Authorization": "Bearer " + pair.other["token"]}).status_code
        == 404
    )
    assert (
        pair.request(
            "student", "POST", url, headers=pair.headers, json={**payload, "text": "Changed duplicate"}
        ).status_code
        == 409
    )


def test_concurrent_peer_grading_uses_one_receipt_and_rejects_changed_answer(pair, monkeypatch):
    import asyncio
    from concurrent.futures import ThreadPoolExecutor
    from course_exchange import ExchangeClient
    import practice_grading

    course = enroll(pair)
    original = practice_grading.practice_grader.grade
    count = []

    async def counted(task, attempt):
        count.append(1)
        await asyncio.sleep(0.05)
        return await original(task, attempt)

    monkeypatch.setattr(practice_grading.practice_grader, "grade", counted)
    client = ExchangeClient(pair.configs["student"])
    body = {
        "actor_id": pair.actor,
        "course_id": "source-course",
        "release_id": course["revision_id"],
        "task_revision_id": "q1",
        "attempt_id": "concurrent-attempt",
        "attempt": {"submitted_answer_payload": {"selected_option_id": "A"}, "result": {"score": 0}},
    }
    with ThreadPoolExecutor(max_workers=4) as pool:
        results = list(pool.map(lambda _: client.call("POST", INTERNAL_PREFIX + "grade", body), range(4)))
    assert len(count) == 1 and len({r["receipt_id"] for r in results}) == 1
    assert all(r["result"]["passed"] is True for r in results)
    from fastapi import HTTPException

    with pytest.raises(HTTPException) as error:
        client.call(
            "POST",
            INTERNAL_PREFIX + "grade",
            {**body, "attempt": {"submitted_answer_payload": {"selected_option_id": "B"}}},
        )
    assert error.value.status_code == 409


def test_hint_and_solution_stay_at_publisher_and_remember_exposure(pair):
    course = enroll(pair)
    from course_exchange import ExchangeClient
    from fastapi import HTTPException

    peer = ExchangeClient(pair.configs["student"])
    body = {
        "actor_id": pair.actor,
        "course_id": "source-course",
        "release_id": course["revision_id"],
        "task_revision_id": "q1",
    }
    with pytest.raises(HTTPException) as error:
        peer.call("POST", INTERNAL_PREFIX + "solution", body)
    assert error.value.status_code == 409
    assert (
        peer.call("POST", INTERNAL_PREFIX + "hint", {**body, "level": 1})["hint"]["content"] == "Add the two numbers."
    )
    base = f"/api/courses/{course['course_id']}/practice"
    attempt = pair.request(
        "student", "POST", base + "/attempts", headers=pair.headers, json={"task_revision_id": "q1"}
    ).json()["attempt"]
    submitted = pair.request(
        "student",
        "POST",
        base + "/attempts/" + attempt["attempt_id"] + "/submit",
        headers=pair.headers,
        json={
            "expected_revision": attempt["revision"],
            "answer_payload": {"selected_option_id": "B"},
            "request_id": "solution-test",
        },
    )
    assert submitted.status_code == 200, submitted.text
    assert peer.call("POST", INTERNAL_PREFIX + "solution", body)["solution"]
    exposed = pair.request(
        "student",
        "POST",
        base + "/attempts/" + attempt["attempt_id"] + "/solution",
        headers=pair.headers,
        json={"expected_revision": submitted.json()["attempt"]["revision"]},
    )
    assert exposed.status_code == 200, exposed.text
    fresh = pair.request(
        "student", "POST", base + "/attempts", headers=pair.headers, json={"task_revision_id": "q1", "resume": False}
    )
    assert fresh.status_code == 200 and fresh.json()["attempt"]["solution_revealed"] and fresh.json()["solution"]


def test_published_ppt_download_binds_course_version_and_grant(pair, monkeypatch):
    from pptx import Presentation
    from io import BytesIO
    import slide_deck_v6_renderer

    def export(_snapshot, destination):
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.notes_slide.notes_text_frame.text = "PRIVATE NOTES"
        deck.save(destination)

    monkeypatch.setattr(slide_deck_v6_renderer, "export_slide_deck_v6_pptx", export)
    course = enroll(pair)
    raw = deepcopy(pair.course)
    raw["teacher_ppt_projects"] = {
        "p1": {
            "title": "Arithmetic slides",
            "status": "ready",
            "confirmed_revision": "m1",
            "manuscript": {"manuscript_revision": "m1"},
            "lesson_ids": ["lecture"],
            "sources": {
                "materials": [],
                "lectures": [{"lesson_id": "lecture", "title": "Addition"}],
                "blocks": [
                    {"block_id": "public", "revision": raw["course_document"]["blocks"][0]["internal_revision"]}
                ],
            },
            "asset_ids": [],
            "last_good_render": {
                "manuscript_revision": "m1",
                "deck": {"pages": []},
                "ppt_manuscript": {"manuscript_revision": "m1"},
            },
        }
    }
    pair.stores["teacher"].save_course_sync("source-course", raw)
    release = pair.publish(ppt_project_ids=["p1"])
    assert pair.publish(ppt_project_ids=["p1"])["release_id"] == release["release_id"]
    pair.request(
        "student",
        "POST",
        f"/api/study/courses/{course['course_id']}/exchange/update",
        headers=pair.headers,
        json={"expected_release_id": course["revision_id"]},
    )
    listing = pair.request("student", "GET", "/api/study/courses", headers=pair.headers).json()[0]
    ppt = listing["published_ppts"][0]
    url = f"/api/study/courses/{course['course_id']}/exchange/ppt/{ppt['asset_id']}/download"
    response = pair.request(
        "student", "POST", url, headers=pair.headers, json={"expected_release_id": listing["revision_id"]}
    )
    assert response.status_code == 200, response.text[:200] if response.status_code != 200 else ""
    assert len(Presentation(BytesIO(response.content)).slides) == 1
    assert (
        pair.request(
            "student", "POST", url, headers=pair.headers, json={"expected_release_id": course["revision_id"]}
        ).status_code
        == 409
    )
    assert (
        pair.request(
            "student",
            "POST",
            url,
            headers={"Authorization": "Bearer " + pair.other["token"]},
            json={"expected_release_id": listing["revision_id"]},
        ).status_code
        == 404
    )


def test_source_import_is_personal_and_revocation_blocks_new_generation(pair):
    course = enroll(pair)
    imported = pair.request(
        "student",
        "POST",
        f"/api/study/courses/{course['course_id']}/sources",
        headers=pair.headers,
        json={"revision_id": course["revision_id"]},
    )
    assert imported.status_code == 200, imported.text
    assert imported.json()["course_origin"]["exchange_origin"]["source_course_id"] == "source-course"
    asset_id = imported.json()["id"]
    from study.repository import StudyRepository

    repo = StudyRepository(pair.stores["student"])
    # Pass the node config explicitly through the same request context as a real
    # student process, without sharing a teacher account or source repository.
    marker = node_context.set("student")
    try:
        repo.require_origin_access(pair.actor, repo.source_origin(pair.actor, asset_id))
        pair.request(
            "teacher",
            "DELETE",
            "/api/teacher/courses/source-course/learners/" + pair.actor,
            headers={"X-User-Id": "teacher"},
        )
        from fastapi import HTTPException

        with pytest.raises(HTTPException) as error:
            repo.require_origin_access(pair.actor, repo.source_origin(pair.actor, asset_id))
        assert error.value.status_code == 404
    finally:
        node_context.reset(marker)
    assert not pair.stores["teacher"].load_data("study_course_sources.json")


def test_teacher_handout_projection_is_published_not_raw_course_body(pair, monkeypatch):
    from teacher_lesson_authoring import TeacherLessonAuthoringRepository
    import teacher_content_projection

    authoring = TeacherLessonAuthoringRepository(
        Path(pair.stores["teacher"]._data_dir) / "teacher_lesson_authoring", canonical_storage=pair.stores["teacher"]
    )
    state = {
        "outline_revision_id": "outline-1",
        "lessons": {
            "lecture": {
                "working_revision_id": "plan-1",
                "working_script_revision_id": "script-1",
                "source_state": "current",
                "revisions": [{"revision_id": "plan-1", "plan": {"sections": []}}],
                "script_revisions": [
                    {
                        "revision_id": "script-1",
                        "publication_eligible": True,
                        "source_lesson_plan_revision_id": "plan-1",
                        "sections": [
                            {
                                "section_node_id": "lecture",
                                "blocks": [
                                    {
                                        "block_id": "real-handout",
                                        "role": "concept",
                                        "content": "Confirmed teacher handout",
                                        "title": "Addition",
                                    }
                                ],
                            }
                        ],
                    }
                ],
            }
        },
    }
    monkeypatch.setattr(authoring, "load", lambda _: deepcopy(state))
    monkeypatch.setattr(teacher_content_projection, "authoring_repository", lambda _: authoring)
    raw = deepcopy(pair.course)
    raw["authoring_surface"] = "teacher"
    pair.stores["teacher"].save_course_sync("source-course", raw)
    course = enroll(pair)
    response = pair.request("student", "GET", f"/api/courses/{course['course_id']}/document", headers=pair.headers)
    assert response.status_code == 200 and "Confirmed teacher handout" in response.text
    assert "2 + 3 = 5." not in response.text and "SECRET" not in response.text


@pytest.mark.parametrize(
    "url",
    [
        "http://example.org",
        "https://user:password@example.org",
        "https://example.org/path",
        "https://example.org/?token=secret",
    ],
)
def test_exchange_rejects_unapproved_plain_http_and_embedded_credentials(url):
    from fastapi import HTTPException

    with pytest.raises(HTTPException):
        ExchangeConfig("student", "student-test", "teacher-test", url, "x" * 32).validate()


def test_companion_goal_summary_is_opt_in_minimal_and_withdrawable_over_real_http(pair,tmp_path,monkeypatch):
    import ai_teacher_state
    from ai_teacher_state import AITeacherRepository
    from companion_learning import CompanionService
    repository=AITeacherRepository(tmp_path/'companion')
    monkeypatch.setattr(ai_teacher_state,'ai_teacher_repository',repository)
    course=enroll(pair);cid=course['course_id'];svc=CompanionService(repository)
    svc.command(pair.actor,cid,action='save_goal',expected_revision=1,request_id='goal-request',value={'text':'考试复习','purpose':'exam'},catalog=[],source_revision='cv')
    svc.command(pair.actor,cid,action='feedback_consent',expected_revision=2,request_id='consent-request',value={'enabled':True},catalog=[],source_revision='cv')
    url=f'/api/study/courses/{cid}/exchange/learning-summary'
    response=pair.request('student','POST',url,headers=pair.headers,json={'request_id':'summary-request'})
    assert response.status_code==200 and response.json()['status']=='delivered',response.text
    analysis=pair.request('teacher','GET','/api/teacher/courses/source-course/exchange/learning-analysis',headers={'X-User-Id':'teacher'}).json()
    assert analysis['roster_count']==1 and analysis['shared_goals'][0]['goal']['text']=='考试复习'
    from companion_teaching_feedback import learning_summary
    pending=pair.repository('student').queue_feedback(pair.actor,cid,request_id='pending-old-summary',kind='learning_summary',payload=learning_summary(svc.read(pair.actor,cid),{}),release_id=course['revision_id'])
    svc.command(pair.actor,cid,action='feedback_consent',expected_revision=3,request_id='withdraw-request',value={'enabled':False},catalog=[],source_revision='cv')
    from fastapi import HTTPException
    with pytest.raises(HTTPException) as stale:
        pair.repository('student').deliver(pair.actor,pending)
    assert stale.value.status_code==409
    withdrawn=pair.request('student','POST',url,headers=pair.headers,json={'request_id':'summary-withdraw'})
    assert withdrawn.status_code==200
    teacher=pair.request('teacher','GET','/api/teacher/courses/source-course/exchange/feedback',headers={'X-User-Id':'teacher'}).json()
    assert '考试复习' not in str(teacher)
