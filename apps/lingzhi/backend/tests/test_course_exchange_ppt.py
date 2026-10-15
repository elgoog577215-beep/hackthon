from zipfile import ZipFile

from fastapi import HTTPException
import pytest
from pptx import Presentation

from course_exchange_ppt import freeze_published_ppts, published_ppt_path, shareable_project, strip_private_ppt_parts
from course_document import CourseBlock, CourseDocument, CourseSection, refresh_document_revision
from storage import Storage


def test_published_ppt_removes_private_notes_and_metadata_and_reopens(tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.notes_slide.notes_text_frame.text = "PRIVATE TEACHER NOTE"
    deck.core_properties.author = "PRIVATE AUTHOR"
    source, clean = tmp_path / "source.pptx", tmp_path / "clean.pptx"
    deck.save(source)
    strip_private_ppt_parts(source, clean)
    with ZipFile(clean) as archive:
        assert not any(
            n.startswith(("ppt/notesSlides/", "ppt/notesMasters/")) or n == "docProps/core.xml"
            for n in archive.namelist()
        )
        assert all(b"PRIVATE" not in archive.read(n) for n in archive.namelist())
    assert len(Presentation(clean).slides) == 1
    assert b"PRIVATE TEACHER NOTE" in ZipFile(source).read("ppt/notesSlides/notesSlide1.xml")


@pytest.fixture
def project(tmp_path):
    doc = refresh_document_revision(
        CourseDocument(
            course_id="c",
            title="Arithmetic",
            sections=[CourseSection(section_id="lecture", title="Addition", position=0)],
            blocks=[CourseBlock(block_id="b", section_id="lecture", position=0, payload={"markdown": "Addition"})],
        )
    ).model_dump(mode="json")
    rendered = {"deck": {"pages": []}, "ppt_manuscript": {"manuscript_revision": "m1"}, "manuscript_revision": "m1"}
    value = {
        "title": "Addition",
        "status": "ready",
        "confirmed_revision": "m1",
        "manuscript": {"manuscript_revision": "m1"},
        "sources": {
            "blocks": [{"block_id": "b", "revision": doc["blocks"][0]["internal_revision"]}],
            "materials": [],
            "lectures": [{"lesson_id": "lecture", "title": "Addition"}],
        },
        "lesson_ids": ["lecture"],
        "asset_ids": [],
        "last_good_render": rendered,
    }
    return Storage(str(tmp_path / "data")), {"teacher_ppt_projects": {"p1": value}}, doc


def test_immutable_export_published_once_and_integrity_checked(project, monkeypatch):
    store, course, doc = project
    import slide_deck_v6_renderer

    calls = []

    def export(_snapshot, destination):
        calls.append(1)
        deck = Presentation()
        deck.slides.add_slide(deck.slide_layouts[6])
        deck.save(destination)

    monkeypatch.setattr(slide_deck_v6_renderer, "export_slide_deck_v6_pptx", export)
    first = freeze_published_ppts(store, course, doc, ["p1"])
    assert first == freeze_published_ppts(store, course, doc, ["p1"]) and len(calls) == 1
    path = published_ppt_path(store, first[0])
    assert len(Presentation(path).slides) == 1
    path.write_bytes(b"corrupt")
    with pytest.raises(HTTPException):
        published_ppt_path(store, first[0])


@pytest.mark.parametrize("change", ["stale_block", "private_block", "uploaded_source", "unconfirmed", "unfinished"])
def test_private_or_stale_ppt_cannot_be_published(project, change):
    store, course, doc = project
    value = course["teacher_ppt_projects"]["p1"]
    if change == "stale_block":
        value["sources"]["blocks"][0]["revision"] = "old"
    elif change == "private_block":
        value["sources"]["blocks"][0]["block_id"] = "teacher-only"
    elif change == "uploaded_source":
        value["asset_ids"] = ["teacher-private-upload"]
    elif change == "unconfirmed":
        value["confirmed_revision"] = "old"
    else:
        value["status"] = "rendering"
    assert not shareable_project(value, doc)
    with pytest.raises(HTTPException) as error:
        freeze_published_ppts(store, course, doc, ["p1"])
    assert error.value.status_code == 409
