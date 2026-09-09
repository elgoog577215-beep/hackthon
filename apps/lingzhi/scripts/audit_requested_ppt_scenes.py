"""Audit all original draft scenes in disposable storage, without task writes."""
import inspect
import json
import os
import sys
import tempfile
from collections import Counter
from pathlib import Path

root = Path("/opt/lingzhi/state/backend-data")
course = "afb29754-6842-437b-af1b-5866bfb53b41"
data = json.loads((root / "teacher_lesson_authoring" / f"{course}.json").read_text(encoding="utf-8"))
manuscript = data["lessons"]["L1-1"]["ppt_manuscript"]["manuscript"]
with tempfile.TemporaryDirectory(prefix="ppt-original-scene-audit-") as directory:
    os.environ["LINGZHI_DATA_DIR"] = str(Path(directory) / "data")
    os.environ["LINGZHI_TASK_RUNTIME_MODE"] = "read_only"
    sys.path.insert(0, "/opt/lingzhi/hackthon/backend")
    from pptx import Presentation
    from pptx.util import Pt
    from ppt_page_scene import ResolvedPageScene
    from ppt_native_scene import render_scene, audit_scene
    from slide_asset_repository import SlideAssetRepository
    from slide_deck_renderer import audit_exported_pptx

    scenes = [ResolvedPageScene.model_validate(s) for p in manuscript["pages"] for s in p["resolved_scenes"]]
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Pt(960), Pt(540)
    assets = SlideAssetRepository(root / "slide_visual_assets")
    for scene in scenes:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        render_scene(slide, scene, assets=assets)
        audit_scene(slide, scene)
    output = Path(directory) / "audit.pptx"
    presentation.save(output)
    before = audit_exported_pptx(output, expected_scenes=scenes, require_pixel_audit=False)
    safe_keys = {"code", "page", "shape_name", "shape_names", "maximum_wrapped_lines", "allowed_wrapped_lines", "minimum_font_size_pt"}
    for issue in before["blockers"]:
        page = int(issue.get("page") or 1)
        obj = next((o for o in scenes[page - 1].objects if "teaching:" + o.object_id == issue.get("shape_name")), None)
        print(json.dumps({"event": "actual_blocker", **{k: v for k, v in issue.items() if k in safe_keys},
                          "scene_slot": obj.slot_id if obj else None}, ensure_ascii=True), flush=True)
    # Evaluate the single proposed classification fix in an isolated function
    # namespace. Deployed source files and course state remain untouched.
    source = inspect.getsource(audit_exported_pptx)
    needle = "and not is_title\n                    and top_inches >= 1.9"
    if source.count(needle) != 1:
        raise ValueError("audit_source_changed")
    source = source.replace(needle, "and not is_title\n                    and not (scene_object is not None and scene_object.slot_id == 'decoration')\n                    and top_inches >= 1.9")
    namespace = dict(audit_exported_pptx.__globals__)
    exec(compile(source, "<isolated-audit-proposal>", "exec"), namespace)
    after = namespace["audit_exported_pptx"](output, expected_scenes=scenes, require_pixel_audit=False)
    print(json.dumps({"event": "all_scenes_audited", "pages": len(scenes),
                      "before": dict(Counter(i["code"] for i in before["blockers"])),
                      "after": dict(Counter(i["code"] for i in after["blockers"])),
                      "proposed_fix_passed": after["passed"]}, ensure_ascii=True), flush=True)
